"""Gera dados/modelo-especies.xlsx a partir das colunas aceitas pelo importador.

Uso (na raiz do projeto):
  npx tsx -e "import {COLUNAS_ESPECIE} from './dominio/src/index'; console.log(JSON.stringify(COLUNAS_ESPECIE))" > /tmp/colunas.json
  python3 dados/gerar-modelo-especies.py /tmp/colunas.json
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

colunas = json.load(open(sys.argv[1], encoding="utf-8"))
FONTE = "Arial"
OBRIGATORIAS = {"nome_cientifico"}
EXPLICACAO = {
    "nome_popular": "Um nome por linha, ou vários separados por / (ex.: Pau-de-óleo/Copaíba).",
    "nome_cientifico": "Obrigatório. Gênero e espécie, com ou sem autor.",
    "meses_flor": "Números (9,10,11), nomes (set out nov) ou intervalo (setembro a novembro).",
    "meses_fruto": "Números (9,10,11), nomes (set out nov) ou intervalo (setembro a novembro).",
    "origem": "nativa, naturalizada, cultivada ou exótica.",
    "lorenzi_volume": "Árvores Brasileiras (Harry Lorenzi): número do volume.",
    "lorenzi_pagina": "Árvores Brasileiras (Harry Lorenzi): página.",
}
EXEMPLO = {
    "nome_popular": "Árvore-de-exemplo",
    "familia": "Ficticiaceae",
    "nome_cientifico": "Ficticia exemplaris Fulano",
    "folha": "composta",
    "disposicao": "alterna",
    "latex": "não",
    "espinhos": "não",
    "casca": "lisa, acinzentada",
    "flor": "amarela, em cachos",
    "fruto": "vagem seca",
    "meses_flor": "set a nov",
    "meses_fruto": "dez a fev",
    "origem": "nativa",
}

wb = Workbook()
ws = wb.active
ws.title = "Espécies"
ws["A1"] = "CRADS · Lista de espécies para importação"
ws["A1"].font = Font(name=FONTE, bold=True, size=14)
ws["A2"] = (
    "Preencha a partir da linha 5. Só a coluna Nome científico é obrigatória; as outras completam a ficha. "
    "A linha 4 é um EXEMPLO FICTÍCIO: apague antes de importar. A ordem das colunas pode mudar; o app reconhece pelo título."
)
ws["A2"].font = Font(name=FONTE, italic=True, size=10)
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=min(len(colunas), 10))
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[2].height = 42

cabecalho = PatternFill("solid", fgColor="1F6F43")
entrada = PatternFill("solid", fgColor="FFF6D5")
for i, (campo, defin) in enumerate(colunas.items(), start=1):
    c = ws.cell(row=3, column=i, value=defin["rotulo"] + (" *" if campo in OBRIGATORIAS else ""))
    c.font = Font(name=FONTE, bold=True, color="FFFFFF")
    c.fill = cabecalho
    c.alignment = Alignment(wrap_text=True, vertical="center")
    if campo in EXPLICACAO:
        c.comment = Comment(EXPLICACAO[campo], "CRADS")
    ex = ws.cell(row=4, column=i, value=EXEMPLO.get(campo))
    ex.font = Font(name=FONTE, italic=True, color="808080")
    ws.column_dimensions[c.column_letter].width = max(14, min(32, len(defin["rotulo"]) + 4))
    for r in range(5, 1005):
        ws.cell(row=r, column=i).font = Font(name=FONTE)
        ws.cell(row=r, column=i).fill = entrada
ws.row_dimensions[3].height = 36
ws.freeze_panes = "A5"

# lista fechada para a origem
idx_origem = list(colunas).index("origem") + 1
letra = ws.cell(row=3, column=idx_origem).column_letter
dv = DataValidation(type="list", formula1='"nativa,naturalizada,cultivada,exótica"', allow_blank=True)
dv.add(f"{letra}5:{letra}1004")
ws.add_data_validation(dv)

leg = wb.create_sheet("Como preencher")
linhas = [
    ("Como preencher", True),
    ("• Células amarelas: para preencher. Cabeçalho verde: não altere o texto.", False),
    ("• Uma linha por nome popular. A mesma espécie pode aparecer em várias linhas: o app junta pelo nome científico.", False),
    ("• Não invente dados: deixe em branco o que não souber. A equipe completa a ficha depois no app.", False),
    ("• O app marca para revisão: famílias diferentes para a mesma espécie, nomes com 'ou', 'et al.' ou só o gênero.", False),
    ("• A importação nunca apaga nem sobrescreve o que já está na ficha: só completa campos vazios e acrescenta nomes populares.", False),
    ("• Também é aceita a planilha de controle atual (aba Espécies com Nome Popular, Família Botânica e Nome Científico).", False),
    ("", False),
    ("Colunas", True),
]
for campo, defin in colunas.items():
    linhas.append((f"{defin['rotulo']}: {EXPLICACAO.get(campo, 'texto livre')}", False))
for r, (texto, negrito) in enumerate(linhas, start=1):
    c = leg.cell(row=r, column=1, value=texto)
    c.font = Font(name=FONTE, bold=negrito, size=12 if negrito else 11)
leg.column_dimensions["A"].width = 120

saida = Path(__file__).with_name("modelo-especies.xlsx")
wb.save(saida)
print(f"gerado {saida}")
